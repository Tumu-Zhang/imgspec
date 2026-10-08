"""输入探测测试：识别类别、页数、源分辨率，并给出可读的错误提示。"""

from __future__ import annotations

from pathlib import Path

import fitz
from PIL import Image

from conftest import make_pdf, make_raster
from imgspec.ingest import SourceKind, classify, probe


class TestRaster:
    def test_basic_probe(self, tmp_path):
        path = make_raster(tmp_path / "fig.png", (1200, 900))
        info = probe(path)
        assert info.kind is SourceKind.RASTER
        assert info.pixel_size == (1200, 900)
        assert info.aspect == 1200 / 900
        assert info.page_count == 1
        assert info.error is None

    def test_reads_embedded_dpi(self, tmp_path):
        path = make_raster(tmp_path / "fig.png", (600, 600), dpi=(300, 300))
        info = probe(path)
        # PNG 用 pHYs 存 DPI（单位像素/米），300dpi 只能近似表示，允许千分位误差
        assert info.dpi is not None
        assert abs(info.dpi[0] - 300.0) < 0.01
        assert "300" in info.source_dpi_text

    def test_missing_dpi_is_flagged(self, tmp_path):
        path = make_raster(tmp_path / "fig.png", (600, 600))
        info = probe(path)
        assert info.dpi is None
        assert any("无 DPI 元数据" in note for note in info.notes)

    def test_alpha_channel_note(self, tmp_path):
        path = tmp_path / "alpha.png"
        Image.new("RGBA", (100, 100), (0, 0, 0, 128)).save(path)
        info = probe(path)
        assert any("透明" in note for note in info.notes)

    def test_zero_dpi_treated_as_missing(self, tmp_path):
        """有些工具会写 dpi=(0, 0)，那不是可用信息，不能被显示成「0 dpi」。"""
        path = tmp_path / "zero.png"
        Image.new("RGB", (100, 100), "white").save(path, dpi=(0, 0))
        info = probe(path)
        assert info.dpi is None
        assert any("无 DPI 元数据" in note for note in info.notes)

    def test_multipage_gif_collapses_to_first_frame(self, tmp_path):
        path = tmp_path / "anim.gif"
        frames = [Image.new("RGB", (40, 40), (i * 40, 0, 0)) for i in range(3)]
        frames[0].save(path, save_all=True, append_images=frames[1:])
        info = probe(path)
        assert info.page_count == 1
        assert any("多帧" in note for note in info.notes)

    def test_tiff_extension(self, tmp_path):
        path = make_raster(tmp_path / "scan.tiff", (300, 300))
        assert probe(path).kind is SourceKind.RASTER


class TestPdf:
    def test_page_count_and_sizes(self, tmp_path):
        path = make_pdf(tmp_path / "doc.pdf", pages=4, size_in=(8.5, 11.0))
        info = probe(path)
        assert info.kind is SourceKind.PDF
        assert info.page_count == 4
        assert len(info.page_sizes_in) == 4
        width_in, height_in = info.page_sizes_in[0]
        assert round(width_in, 3) == 8.5
        assert round(height_in, 3) == 11.0

    def test_mixed_page_sizes_warn(self, tmp_path):
        path = tmp_path / "mixed.pdf"
        doc = fitz.open()
        doc.new_page(width=612, height=792)
        doc.new_page(width=400, height=400)
        doc.save(path)
        doc.close()
        info = probe(path)
        assert any("尺寸不一致" in note for note in info.notes)

    def test_corrupt_pdf_reports_error(self, tmp_path):
        path = tmp_path / "broken.pdf"
        path.write_bytes(b"%PDF-1.4\nnot really a pdf")
        info = probe(path)
        assert info.error is not None


class TestSlides:
    def test_pptx_probe(self, tmp_path):
        from pptx import Presentation
        from pptx.util import Inches

        path = tmp_path / "deck.pptx"
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank = prs.slide_layouts[6]
        for _ in range(3):
            prs.slides.add_slide(blank)
        prs.save(path)

        info = probe(path)
        assert info.kind is SourceKind.SLIDES
        assert info.page_count == 3
        assert info.aspect is not None and abs(info.aspect - 13.333 / 7.5) < 0.01

    def test_legacy_ppt_defers_to_powerpoint(self, tmp_path):
        path = tmp_path / "old.ppt"
        path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1stub")
        info = probe(path)
        assert info.kind is SourceKind.SLIDES
        assert any("PowerPoint" in note for note in info.notes)


class TestFailures:
    def test_missing_file(self, tmp_path):
        info = probe(tmp_path / "nope.png")
        assert info.error == "文件不存在"

    def test_unsupported_with_hint(self, tmp_path):
        path = tmp_path / "drawing.cdr"
        path.write_bytes(b"RIFFxxxxCDR")
        info = probe(path)
        assert info.kind is SourceKind.UNSUPPORTED
        assert "CorelDRAW" in info.error

    def test_svg_is_supported(self, tmp_path):
        """SVG 由 MuPDF 直接渲染，不再需要外部工具。"""
        path = tmp_path / "figure.svg"
        path.write_text(
            '<svg xmlns="http://www.w3.org/2000/svg" width="400" height="300">'
            '<rect width="400" height="300" fill="#fff"/></svg>',
            encoding="utf-8",
        )
        info = probe(path)
        assert info.kind is SourceKind.SVG
        assert info.error is None
        assert info.aspect == 400 / 300
        assert info.kind_label == "矢量图"

    def test_eps_hint_mentions_ghostscript(self, tmp_path):
        path = tmp_path / "old.eps"
        path.write_bytes(b"%!PS")
        info = probe(path)
        assert "Ghostscript" in info.error

    def test_unknown_extension(self, tmp_path):
        path = tmp_path / "data.xyz"
        path.write_text("hello", encoding="utf-8")
        info = probe(path)
        assert info.kind is SourceKind.UNSUPPORTED

    def test_corrupt_image_reports_error(self, tmp_path):
        path = tmp_path / "fake.png"
        path.write_bytes(b"this is not a png")
        info = probe(path)
        assert info.error is not None


class TestClassify:
    def test_categories(self):
        assert classify(Path("a.PDF")) is SourceKind.PDF
        assert classify(Path("a.pptx")) is SourceKind.SLIDES
        assert classify(Path("a.PPT")) is SourceKind.SLIDES
        assert classify(Path("a.jpeg")) is SourceKind.RASTER
        assert classify(Path("a.TIF")) is SourceKind.RASTER
        assert classify(Path("a.docx")) is SourceKind.UNSUPPORTED
